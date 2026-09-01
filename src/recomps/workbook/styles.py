"""Shared workbook styling.

One typeface, one set of fills, defined once so five sheets cannot drift apart.
The convention that matters to a reader: **blue text on yellow is an input you
may edit**, and every stat that depends on it recalculates. Everything else is
either data pulled from a source or a formula over that data.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"

TITLE = Font(name=FONT, size=14, bold=True)
SUBTITLE = Font(name=FONT, size=9, italic=True, color="595959")
HEADER = Font(name=FONT, size=10, bold=True, color="FFFFFF")
SECTION = Font(name=FONT, size=11, bold=True)
BODY = Font(name=FONT, size=10)
BODY_BOLD = Font(name=FONT, size=10, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="595959")
EDITABLE = Font(name=FONT, size=10, color="0000CC", bold=True)
PRIMARY = Font(name=FONT, size=10, bold=True, color="1F4E79")

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
EDITABLE_FILL = PatternFill("solid", fgColor="FFF2A8")
PRIMARY_FILL = PatternFill("solid", fgColor="DDEBF7")
SHORTLIST_FILL = PatternFill("solid", fgColor="D8EFD5")
CAUTION_FILL = PatternFill("solid", fgColor="F8D7DA")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center")


def style_header(cell) -> None:
    cell.font = HEADER
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_editable(cell, number_format: str | None = None) -> None:
    cell.font = EDITABLE
    cell.fill = EDITABLE_FILL
    cell.border = BOX
    if number_format:
        cell.number_format = number_format
