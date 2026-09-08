"""Saved searches, saved runs, and reopening them.

A profile is a saved question; a run is one answer to it. These cover the parts
that let you pick a question you asked before, look at the answer you got, and
ask a variation of it without going back to any website.
"""

from __future__ import annotations

from datetime import UTC, date, datetime, timedelta

import pytest

from recomps.config.profile import Subject, vacant_land_profile
from recomps.config.store import (
    APP_DIR_ENV,
    available_profiles,
    delete_user_profile,
    describe_profile,
    load_user_profiles,
    profile_origin,
    resolve_profile,
    save_user_profile,
)
from recomps.markets.demoville import MARKET
from recomps.model.snapshot import Snapshot, build_snapshot
from recomps.pipeline import reopen as reopen_mod
from recomps.pipeline.run import run_pipeline
from recomps.reporting import history as history_mod
from recomps.research.fixture import FixtureResearcher
from recomps.workbook.builder import build_workbook

AS_OF = date(2026, 8, 31)


@pytest.fixture
def config_dir(tmp_path, monkeypatch):
    """Point profile storage at a temp directory so tests never touch real config."""
    monkeypatch.setenv(APP_DIR_ENV, str(tmp_path / "config"))
    return tmp_path / "config"


@pytest.fixture(scope="module")
def result():
    return run_pipeline(
        MARKET, MARKET.profiles()["demo-lots"], FixtureResearcher(), window_end=AS_OF
    )


def _profile(name="my-lot", size=6450.0):
    profile = vacant_land_profile(name)
    profile.subject = Subject(lot_sqft=size, label="Your lot")
    return profile


# ---------------------------------------------------------------------------
# Saving and picking profiles
# ---------------------------------------------------------------------------


def test_a_saved_profile_survives_and_reloads(config_dir):
    save_user_profile("demoville", _profile())
    reloaded = load_user_profiles("demoville")
    assert "my-lot" in reloaded
    assert reloaded["my-lot"].subject.lot_sqft == 6450.0


def test_saved_profiles_appear_alongside_the_market_s_own(config_dir):
    save_user_profile("demoville", _profile())
    names = set(available_profiles(MARKET))
    assert {"demo-lots", "demo-sfr", "my-lot"} <= names


def test_a_saved_profile_shadows_a_market_one_of_the_same_name(config_dir):
    """The market describes the place; you describe your property."""
    mine = _profile("demo-lots", size=9999.0)
    save_user_profile("demoville", mine)
    assert resolve_profile(MARKET, "demo-lots").subject.lot_sqft == 9999.0
    assert profile_origin(MARKET, "demo-lots") == "yours"


def test_market_profiles_are_labelled_as_such(config_dir):
    assert profile_origin(MARKET, "demo-sfr") == "market"


def test_deleting_a_saved_profile(config_dir):
    save_user_profile("demoville", _profile())
    assert delete_user_profile("demoville", "my-lot") is True
    assert "my-lot" not in load_user_profiles("demoville")
    assert delete_user_profile("demoville", "my-lot") is False


def test_deleting_a_shadowing_profile_reveals_the_market_s_version(config_dir):
    save_user_profile("demoville", _profile("demo-lots", size=9999.0))
    delete_user_profile("demoville", "demo-lots")
    assert resolve_profile(MARKET, "demo-lots").subject.lot_sqft == 6450.0


def test_unknown_profile_names_the_alternatives(config_dir):
    with pytest.raises(KeyError) as excinfo:
        resolve_profile(MARKET, "nope")
    assert "demo-lots" in str(excinfo.value)


def test_the_one_line_description_says_what_it_searches_for():
    text = describe_profile(MARKET.profiles()["demo-lots"])
    assert "vacant land" in text
    assert "6,450" in text
    assert "5,000-8,000" in text


# ---------------------------------------------------------------------------
# Archiving runs under their profile
# ---------------------------------------------------------------------------


def _archive(tmp_path, result, when: datetime):
    directory = history_mod.run_dir(
        tmp_path, result.market_name, result.profile.name, when
    )
    directory.mkdir(parents=True, exist_ok=True)
    snapshot = build_snapshot(result)
    snapshot.payload["run"]["run_at"] = when.isoformat()
    snapshot.write(directory / history_mod.SNAPSHOT_NAME)
    build_workbook(result).save(directory / history_mod.WORKBOOK_NAME)
    return directory


def test_runs_are_grouped_by_profile(tmp_path, result):
    base = datetime(2026, 8, 4, tzinfo=UTC)
    _archive(tmp_path, result, base)
    _archive(tmp_path, result, base + timedelta(days=14))

    records = history_mod.list_runs(tmp_path, "demoville", "demo-lots")
    assert len(records) == 2
    assert records[0].run_at > records[1].run_at, "newest first"
    assert all(r.profile == "demo-lots" for r in records)


def test_history_can_be_narrowed_to_one_profile(tmp_path, result):
    _archive(tmp_path, result, datetime(2026, 8, 4, tzinfo=UTC))
    assert history_mod.list_runs(tmp_path, "demoville", "demo-lots")
    assert history_mod.list_runs(tmp_path, "demoville", "demo-sfr") == []


def test_the_archive_keeps_the_workbook_not_just_the_data(tmp_path, result):
    directory = _archive(tmp_path, result, datetime(2026, 8, 4, tzinfo=UTC))
    record = history_mod.list_runs(tmp_path)[0]
    assert record.workbook is not None and record.workbook.stat().st_size > 0
    assert record.snapshot is not None
    assert record.directory == directory


def test_latest_run_excludes_the_one_being_written(tmp_path, result):
    """--compare last means the run before this one, never this one."""
    first = datetime(2026, 8, 4, tzinfo=UTC)
    second = datetime(2026, 8, 20, tzinfo=UTC)
    _archive(tmp_path, result, first)
    _archive(tmp_path, result, second)

    previous = history_mod.latest_run(tmp_path, "demoville", "demo-lots", before=second)
    assert previous is not None
    assert previous.run_at == first


def test_latest_run_compares_across_naive_and_aware_timestamps(tmp_path, result):
    """Folder stamps are UTC; a caller may pass a naive datetime."""
    _archive(tmp_path, result, datetime(2026, 8, 4, tzinfo=UTC))
    naive = datetime(2026, 8, 20)
    assert history_mod.latest_run(tmp_path, "demoville", "demo-lots", before=naive)


def test_no_runs_yet_is_not_an_error(tmp_path):
    assert history_mod.list_runs(tmp_path) == []
    assert history_mod.latest_run(tmp_path, "demoville", "demo-lots") is None


def test_profile_names_with_awkward_characters_are_still_usable(tmp_path, result):
    directory = history_mod.run_dir(
        tmp_path, "demoville", "my lot / 2026", datetime(2026, 8, 4, tzinfo=UTC)
    )
    directory.mkdir(parents=True)
    assert "/" not in directory.parent.name
    assert history_mod.list_runs(tmp_path)[0].profile == history_mod.safe_name("my lot / 2026")


# ---------------------------------------------------------------------------
# Reopening a saved run
# ---------------------------------------------------------------------------


def test_reopening_reproduces_the_original_numbers(result):
    reopened = reopen_mod.reopen(build_snapshot(result), MARKET)
    assert reopened.sold_stats.count == result.sold_stats.count
    assert reopened.sold_stats.median_ppsf == pytest.approx(result.sold_stats.median_ppsf)
    assert reopened.valuation.primary.value == pytest.approx(result.valuation.primary.value)
    assert reopened.window_start == result.window_start
    assert reopened.window_end == result.window_end


def test_reopening_uses_the_profile_the_run_actually_used(result):
    """A profile edited since the run must not change what the run meant."""
    snapshot = build_snapshot(result)
    restored = reopen_mod.profile_from_snapshot(snapshot)
    assert restored.name == result.profile.name
    assert restored.subject_size() == result.profile.subject_size()
    assert restored.window_days == result.profile.window_days


def test_a_regenerated_workbook_is_valid(result, tmp_path):
    reopened = reopen_mod.reopen(build_snapshot(result), MARKET)
    path = tmp_path / "again.xlsx"
    build_workbook(reopened).save(path)
    assert path.stat().st_size > 0

    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb.sheetnames[0] == "Summary"
    assert wb["Sold Comps"].max_row >= reopened.sold_stats.count + 1


def test_reopening_with_a_different_subject_size_moves_the_valuation(result):
    from recomps.config.profile import CompProfile

    bigger = CompProfile.from_dict(result.profile.to_dict())
    bigger.subject.lot_sqft = 9000.0
    reopened = reopen_mod.reopen(build_snapshot(result), MARKET, profile=bigger)

    assert reopened.valuation.primary.value > result.valuation.primary.value
    # Same rate, different size: the sales did not change, the question did.
    assert reopened.valuation.primary.ppsf == pytest.approx(result.valuation.primary.ppsf)


def test_dropping_a_comp_changes_the_answer_and_says_so(result):
    victim = result.sold[0].address
    reopened = reopen_mod.reopen(
        build_snapshot(result), MARKET, exclude_addresses=[victim]
    )
    assert reopened.sold_stats.count == result.sold_stats.count - 1
    assert victim not in {c.address for c in reopened.sold}
    assert any("Reopened from a saved run" in c for c in reopened.caveats)


def test_an_unadjusted_reopen_adds_no_caveat(result):
    reopened = reopen_mod.reopen(build_snapshot(result), MARKET)
    assert not any("Reopened from a saved run" in c for c in reopened.caveats)


def test_reopening_survives_a_round_trip_through_disk(result, tmp_path):
    path = build_snapshot(result).write(tmp_path / "snap.json")
    reopened = reopen_mod.reopen_path(str(path), MARKET)
    assert reopened.sold_stats.count == result.sold_stats.count
    assert reopened.agent_analysis.attributed == result.agent_analysis.attributed


def test_a_snapshot_without_run_metadata_is_rejected():
    broken = Snapshot(schema_version=1, payload={"sold": [], "active": []})
    with pytest.raises(reopen_mod.SnapshotUnreadable):
        reopen_mod.reopen(broken, MARKET)


# ---------------------------------------------------------------------------
# Nothing can go missing from a saved run
# ---------------------------------------------------------------------------


def test_reopening_keeps_every_single_row(result):
    """The worry this answers: can a comp fall off when you reopen a past run?

    It cannot. The saved run holds every row it collected, reopening reads all
    of them, and nothing goes back to any website.
    """
    reopened = reopen_mod.reopen(build_snapshot(result), MARKET)

    assert [c.address for c in reopened.sold] == [c.address for c in result.sold]
    assert [c.address for c in reopened.active] == [c.address for c in result.active]
    for before, after in zip(result.sold, reopened.sold, strict=True):
        assert after.sold_price == before.sold_price
        assert after.lot_sqft == before.lot_sqft
        assert after.sold_date == before.sold_date
        assert after.brokerage == before.brokerage
        assert after.agent == before.agent
        assert after.final_list_price == before.final_list_price
        assert after.original_list_price == before.original_list_price


def test_reopening_keeps_rows_that_could_not_contribute_a_rate(result):
    """A sale with no published size is still a sale, and must survive."""
    sizeless = [c.address for c in result.sold if c.lot_sqft is None]
    assert sizeless, "the fixture set should contain one, or this proves nothing"
    reopened = reopen_mod.reopen(build_snapshot(result), MARKET)
    assert set(sizeless) <= {c.address for c in reopened.sold}


def test_viewing_a_past_run_recomputes_nothing(result):
    """`read_stored` hands back the figures the run itself reported."""
    snapshot = build_snapshot(result)
    stored = reopen_mod.read_stored(snapshot)

    assert stored.sold_count == result.sold_stats.count
    assert stored.primary_value == result.valuation.primary.value
    assert stored.guidance["floor"] == result.guidance.floor
    assert len(stored.sold) == len(result.sold)
    assert stored.profile_name == result.profile.name
    assert stored.window_start == result.window_start.isoformat()


def test_the_stored_view_and_a_reopen_agree_today(result):
    """They should match now. If they ever diverge, the analysis has changed --
    which is the point of keeping both: the stored view is the historical
    record, the reopened view is what today's code makes of the same rows."""
    snapshot = build_snapshot(result)
    stored = reopen_mod.read_stored(snapshot)
    reopened = reopen_mod.reopen(snapshot, MARKET)
    assert stored.primary_value == pytest.approx(reopened.valuation.primary.value)
    assert stored.sold_count == reopened.sold_stats.count


def test_a_stored_view_needs_no_market_plugin(result, tmp_path):
    """Viewing history must work even if the market definition is unavailable."""
    path = build_snapshot(result).write(tmp_path / "snap.json")
    stored = reopen_mod.read_stored_path(str(path))
    assert stored.primary_value and stored.sold


# ---------------------------------------------------------------------------
# Local time on screen, UTC on disk
# ---------------------------------------------------------------------------
#
# A run started at half past ten on a Friday evening in California is already
# Saturday in UTC. Labelling it Saturday tells the reader something false about
# their own week, and the date is what they use to find the run again.


def test_a_run_is_labelled_in_local_time():
    from pathlib import Path

    from recomps.clock import to_local

    when = datetime(2026, 9, 5, 5, 30, tzinfo=UTC)
    record = history_mod.RunRecord(
        market="m", profile="p", run_at=when, directory=Path("x")
    )
    # Twelve-hour, and the hour is not zero-padded: "1:40 PM", not "01:40 PM".
    assert record.label == to_local(when).strftime("%Y-%m-%d %I:%M %p").replace(
        " 0", " ", 1
    )
    assert "AM" in record.label or "PM" in record.label
    assert record.local_date == to_local(when).strftime("%Y-%m-%d")


def test_the_archive_directory_stays_utc():
    """Local stamps do not sort across a daylight-saving boundary, and a run
    has to mean the same thing on another machine."""
    when = datetime(2026, 9, 5, 5, 30, tzinfo=UTC)
    directory = history_mod.run_dir("/data", "m", "p", when)
    assert directory.name == "20260905T053000Z"


def test_a_run_is_found_by_the_date_it_is_listed_under(tmp_path, result):
    """`--run <date>` has to match what history printed, not the UTC date."""
    when = datetime(2026, 9, 5, 5, 30, tzinfo=UTC)
    directory = history_mod.run_dir(tmp_path, "demoville", "demo-lots", when)
    directory.mkdir(parents=True)
    build_snapshot(result).write(directory / history_mod.SNAPSHOT_NAME)

    record = history_mod.list_runs(tmp_path)[0]
    assert record.local_date == record.label[:10]
    assert record.directory.name.startswith("20260905")


def test_the_workbook_dates_itself_by_the_local_day(result):
    from recomps.clock import local_date

    assert result.pull_date == local_date(result.run_at)


def test_a_naive_stamp_is_read_as_utc_not_as_local():
    """Everything this project stores is UTC. Guessing local for a stored stamp
    would silently shift every archived run by the offset."""
    from recomps.clock import to_local

    naive = datetime(2026, 9, 5, 5, 30)
    aware = datetime(2026, 9, 5, 5, 30, tzinfo=UTC)
    assert to_local(naive) == to_local(aware)


def test_an_unparseable_stamp_does_not_break_a_label():
    from recomps.model.snapshot import Snapshot

    snapshot = Snapshot(
        schema_version=1, payload={"run": {"market": "m", "run_at": "not a date"}}
    )
    assert snapshot.local_date == "not a date"[:10]


# ---------------------------------------------------------------------------
# Renaming a search takes its past runs with it
# ---------------------------------------------------------------------------


def test_renaming_moves_the_archived_runs_too(config_dir, tmp_path, result):
    """The name is part of the path a run was saved under. A rename that only
    touched the profile would orphan every run of it."""
    from recomps.config.store import rename_user_profile, save_user_profile

    save_user_profile("demoville", _profile("old-name"))
    when = datetime(2026, 8, 31, tzinfo=UTC)
    directory = history_mod.run_dir(tmp_path, "demoville", "old-name", when)
    directory.mkdir(parents=True)
    build_snapshot(result).write(directory / history_mod.SNAPSHOT_NAME)

    moved = history_mod.rename_profile_runs(tmp_path, "demoville", "old-name", "new-name")
    assert moved == 1
    assert rename_user_profile("demoville", "old-name", "new-name")

    assert not history_mod.list_runs(tmp_path, "demoville", "old-name")
    kept = history_mod.list_runs(tmp_path, "demoville", "new-name")
    assert len(kept) == 1 and kept[0].snapshot is not None
    saved = load_user_profiles("demoville")
    assert "new-name" in saved and "old-name" not in saved
    assert saved["new-name"].name == "new-name"


def test_a_rename_onto_an_existing_name_is_refused(config_dir, tmp_path, result):
    from recomps.config.store import rename_user_profile, save_user_profile

    save_user_profile("demoville", _profile("one"))
    save_user_profile("demoville", _profile("two"))
    with pytest.raises(ValueError):
        rename_user_profile("demoville", "one", "two")
    assert set(load_user_profiles("demoville")) == {"one", "two"}


def test_a_rename_onto_an_existing_archive_is_refused(tmp_path, result):
    when = datetime(2026, 8, 31, tzinfo=UTC)
    for name in ("one", "two"):
        d = history_mod.run_dir(tmp_path, "demoville", name, when)
        d.mkdir(parents=True)
        build_snapshot(result).write(d / history_mod.SNAPSHOT_NAME)
    with pytest.raises(FileExistsError):
        history_mod.rename_profile_runs(tmp_path, "demoville", "one", "two")
    assert history_mod.list_runs(tmp_path, "demoville", "one")


def test_renaming_a_search_with_no_runs_is_fine(config_dir, tmp_path):
    from recomps.config.store import rename_user_profile, save_user_profile

    save_user_profile("demoville", _profile("lonely"))
    assert history_mod.rename_profile_runs(tmp_path, "demoville", "lonely", "renamed") == 0
    assert rename_user_profile("demoville", "lonely", "renamed")


def test_a_market_s_own_search_cannot_be_renamed(config_dir):
    """It belongs to the market definition; renaming here would leave the
    original in place and a copy beside it."""
    from recomps.config.store import rename_user_profile

    assert rename_user_profile("demoville", "demo-lots", "mine") is False


def test_the_hour_is_unpadded_without_mangling_the_date():
    """Stripping the leading zero from the result instead of the format turned
    05:30 on the 5th into "2026-09-5 05:30 AM" — the first "05" is in the date.
    It passed in one timezone and failed in another, which is why CI found it
    and a laptop did not."""
    from recomps.clock import local_stamp, to_local

    for moment in (
        datetime(2026, 9, 5, 5, 30, tzinfo=UTC),
        datetime(2026, 9, 5, 12, 30, tzinfo=UTC),
        datetime(2026, 10, 10, 17, 10, tzinfo=UTC),
        datetime(2026, 1, 1, 0, 5, tzinfo=UTC),
    ):
        stamp = local_stamp(moment)
        local = to_local(moment)
        assert stamp.startswith(local.strftime("%Y-%m-%d")), stamp
        assert stamp.endswith(("AM", "PM")), stamp
        hour = stamp[11:].split(":")[0]
        assert not hour.startswith("0"), stamp
        assert hour == local.strftime("%I").lstrip("0"), stamp
